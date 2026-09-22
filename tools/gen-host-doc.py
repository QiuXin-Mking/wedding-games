# -*- coding: utf-8 -*-
"""生成《主持人培训手册》.docx —— 全篇只用主持人视角，不出现任何技术细节。"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

SEAL = RGBColor(0xC8, 0x36, 0x2F)
GOLD = RGBColor(0xA8, 0x85, 0x20)
INK = RGBColor(0x3D, 0x2B, 0x23)
INK2 = RGBColor(0x6B, 0x53, 0x44)

doc = Document()

# 中文字体
style = doc.styles["Normal"]
style.font.name = "微软雅黑"
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
style.paragraph_format.line_spacing = 1.5
style.paragraph_format.space_after = Pt(6)

for sec in doc.sections:
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.4)


def _cn(run, font="微软雅黑"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    return run


def title(text, sub=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = _cn(p.add_run(text), "宋体")
    r.font.size = Pt(26)
    r.bold = True
    r.font.color.rgb = INK
    if sub:
        q = doc.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        s = _cn(q.add_run(sub))
        s.font.size = Pt(10.5)
        s.font.color.rgb = INK2


def h1(text):
    doc.add_paragraph()
    p = doc.add_paragraph()
    r = _cn(p.add_run(text), "微软雅黑")
    r.font.size = Pt(15)
    r.bold = True
    r.font.color.rgb = SEAL
    p.paragraph_format.space_after = Pt(4)


def h2(text):
    p = doc.add_paragraph()
    r = _cn(p.add_run(text))
    r.font.size = Pt(11.5)
    r.bold = True
    r.font.color.rgb = INK
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)


def para(text, bold_head=None):
    p = doc.add_paragraph()
    if bold_head:
        _cn(p.add_run(bold_head)).bold = True
    _cn(p.add_run(text))
    return p


def bullet(text, head=None):
    p = doc.add_paragraph(style="List Bullet")
    if head:
        r = _cn(p.add_run(head))
        r.bold = True
    _cn(p.add_run(text))
    p.paragraph_format.space_after = Pt(2)
    return p


def numbered(text, head=None):
    p = doc.add_paragraph(style="List Number")
    if head:
        _cn(p.add_run(head)).bold = True
    _cn(p.add_run(text))
    p.paragraph_format.space_after = Pt(2)
    return p


def say(text):
    """可以直接念出口的话"""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    r = _cn(p.add_run("〖照着念〗 "))
    r.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = GOLD
    r2 = _cn(p.add_run(text), "楷体")
    r2.font.size = Pt(11)
    r2.font.color.rgb = INK
    return p


def warn(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    r = _cn(p.add_run("⚠ "))
    r.font.color.rgb = SEAL
    r.bold = True
    r2 = _cn(p.add_run(text))
    r2.font.color.rgb = SEAL
    r2.bold = True
    return p


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        r = _cn(cell.paragraphs[0].add_run(h))
        r.bold = True
        r.font.size = Pt(10)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            for seg, bold in v if isinstance(v, list) else [(v, False)]:
                r = _cn(p.add_run(seg))
                r.font.size = Pt(9.5)
                r.bold = bold
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


# ══════════════════════════════════════════════════
title("婚礼互动答题 · 主持人手册", "2026 年 10 月 6 日 ｜ 全场 13 道题，约 15 分钟")

h1("一、先把这件事讲清楚")

para("宾客扫桌上的二维码进入，系统给每人起一个名字（比如「描金的折扇」「晶莹的冰棍」），"
     "不用打字、不用登录。你念题，大屏和每个人的手机同时出现这道题和倒计时。"
     "答得又快又对分越高。13 道题答完，大屏打出最终排名。")

h2("你的位置")
para("你是这场游戏唯一的发动机。系统不会自己往下走 —— "
     "每一道题都要你按「下一题」才会开始。这是故意的："
     "什么时候念完、什么时候观众准备好了，只有你知道。")

h2("三个屏幕，各看各的")
table(
    ["谁", "看到什么", "你要管吗"],
    [
        ["宾客手机", "自己的名字、当前题目、四个选项、自己的得分", "不用，他们自己点"],
        ["大屏幕", "题目、倒计时、每题结果、排行榜", "不用，会自动跟着走"],
        [[("你的控制台", True)], "题目和正确答案、已答人数、所有按钮", "这是你唯一要操作的东西"],
    ],
    [2.6, 6.4, 5.2],
)
warn("正确答案只出现在你的控制台上。别把你的手机屏幕转向观众，也别投屏。")

# ──────────────────────────────────────────────
h1("二、开场前十分钟")

numbered("在你的手机或平板上打开控制台链接（婚礼当天会有人帮你打开并交到你手上）。", "打开控制台。")
numbered("看一眼大屏左上角的小角标。它应该写着「正式题库」。"
         "如果写的是「测试题库」，立刻找技术的人 —— 那套题跟新人一点关系都没有。", "确认题库。")
numbered("桌上的二维码卡片是不是每桌都有。没有的桌子，宾客进不来。", "扫一眼桌卡。")
numbered("控制台上会显示「已入场 N 人」。开场前它应该是 0 或很小的数字。", "看人数。")

warn("大屏如果显示着地址栏，请负责大屏的人点一下「全屏」。地址栏里有口令，露出来观众就能进你的控制台。")

# ──────────────────────────────────────────────
h1("三、开场")

h2("第 1 步 · 喊扫码，然后垫大约 15 秒")
para("三四百人同时扫码，需要一点时间。宾客一秒内就能看到「正在进入」，不会白屏，"
     "但你要给他们把名字看清楚的时间。")
say("桌上都有一张二维码，微信扫一扫就行，不用打字。扫完系统会给各位起个名字，"
    "等下上大屏就认这个名字。名字是随机的，抽到什么算什么 —— "
    "待会儿大屏上要是出现「会响的拨浪鼓」拿了第一名，那就是我们在座的某一位。")

h2("第 2 步 · 看人数，够了就开始")
para("控制台上的「已入场」会一直往上跳。不用等所有人 —— 迟到的人中途也能进来，"
     "进来就能从下一题开始答。")
say("看来大家都进来了。我们开始第一题。")
para("点「开始游戏」。")

# ──────────────────────────────────────────────
h1("四、每一道题：固定的四个动作")

para("这四步从第 1 题到第 13 题，一模一样。记住顺序，全场就不会乱。")

h2("① 先念题，再按按钮")
para("把题目和四个选项完整念一遍。念的时候倒计时还没开始，不占宾客的时间。")
warn("顺序绝对不能反。按下「下一题」的那一瞬间，20 秒倒计时立刻开始走。"
     "如果你先按了再念，宾客会一边听你念一边看着时间掉。")

h2("② 按「下一题」，然后等")
para("题目同时出现在大屏和所有人手机上。这 20 秒你不用做任何事，"
     "可以看着控制台上的「已作答 N 人」往上涨。")
say("（倒计时过半时）还有十秒 —— 没点的抓紧，手机上直接点选项就行。")

h2("③ 唱分")
para("时间一到自动结算。你的控制台上会出现四个数字，这就是你唱分的全部素材：")
table(
    ["控制台显示", "怎么用"],
    [
        ["这题有多少人作答", "人气。数字大就说「几乎全场都参与了」"],
        ["其中多少人答对", "悬念。答对的人少，这题就值得停下来说两句"],
        ["正确答案是哪个", "公布。大屏上同时会高亮出来"],
        ["目前第一名是谁、多少分", "把名字念出来，这是全场最兴奋的一句"],
    ],
    [4.6, 9.6],
)
say("这一题，八十六位来宾作答，只有二十三位答对 —— 正确答案是 B！"
    "现在排在第一名的是「描金的折扇」，两千零三十分。这位朋友，请举个手让大家看看。")

h2("④ 按「下一题」，进入下一轮")
para("回到第 ① 步。")

h2("关于分数，被问到时这么解释")
para("答对得分，答得越快分越高 —— 同样答对，第 2 秒点的人比第 18 秒点的人分高得多。"
     "答错和没答都是 0 分。所以「又快又准」才能拿第一，光靠蒙不行。")

# ──────────────────────────────────────────────
h1("五、六个救场按钮")

para("这些按钮收在控制台的「救场操作」里，平时用不上。"
     "遇到状况时，先看这张表，再决定按哪个。")

table(
    ["按钮", "什么时候用", "按下去会发生什么"],
    [
        ["延长 10 秒", "大家还在低头找选项，明显没答完",
         [("给 10 秒看清题的机会。", False), ("但这 10 秒里提交的一律 0 分", True), ("，"
          "所以这是「让大家看完」，不是「让大家得分」。", False)]],
        ["提前结算本题", "该答的都答了，不想全场干等剩下的十几秒",
         "立刻结算。没提交的记作未作答。"],
        ["重新公布结果", "大屏卡住了，结果没显示出来",
         "只是把画面重播一遍，分数一分不动。这个按钮很安全。"],
        ["把二维码打回大屏", "有人迟到，没赶上扫码",
         "大屏弹出一个大二维码，25 秒后自动收回，游戏不受影响。"],
        [[("跳过本题", True), ("（危险）", False)], "题目本身有问题，不能算数",
         [("全场这一题都是 0 分，", False), ("按下去撤不回来", True),
          ("。大屏显示「本题作废」，不公布答案。", False)]],
        [[("换一道备用题", True), ("（危险）", False)], "你念错了题，或者题目有歧义",
         [("原题作废，换一道新题重来。", False), ("原来那题不能重答", True),
          (" —— 答案刚公布过，重答等于给全场送分。", False)]],
    ],
    [3.0, 4.4, 6.8],
)
para("两个危险操作会弹出确认框，问你一次。其余的按了就执行。")

h2("最常见的一种情况")
para("第一题往往最慢 —— 大家还在适应。如果第一题结束时作答人数明显偏低，别慌，这是正常的。")
say("第一题大家还在找手感，没关系。从第二题开始，题目一出现就在你手机上，直接点就行。")

# ──────────────────────────────────────────────
h1("六、发奖：怎么把人找出来")

para("这是最容易冷场的环节，请特别留意。")
para("名字是系统随机起的，和真人没有任何联系。你念「今天很开心的柯基」，"
     "台下可能有人听成了别的，也可能本人不敢举手 —— 怕举错了更丢脸。")

h2("所以不要只靠念")
para("在控制台的「宾客管理」里搜到这个名字，点一下「点名」。")
para("那位宾客的手机会立刻全屏亮起、并且震动，上面是一行大字：「叫的就是您，请举手」。"
     "其他人的手机完全不受影响。")
say("第一名是「描金的折扇」—— 我已经让系统给这位朋友的手机发了信号，"
    "现在您的手机应该亮起来了，请举个手！")
para("这样他会非常确定叫的是自己，会很干脆地举手。全场的注意力也跟着他走。")

h2("结束后")
para("点「导出明细 CSV」，会存下一份完整成绩单，用 Excel 打开中文不会乱码。")

# ──────────────────────────────────────────────
h1("七、出意外了怎么办")

para("按严重程度分四级。每一级都配好了可以直接念出口的话 —— 慌的时候不要临场编。")

h2("一级 · 大屏黑了")
para("先记住一件事：", )
warn("题目在每一个人的手机上都有。大屏黑了游戏照样能继续，不用暂停。")
say("大屏这边稍微调整一下，大家看手机就可以，我们继续。")
para("同时让负责大屏的人刷新一下页面，会自动回到当前这道题。刷新之后记得重新点「全屏」。")

h2("二级 · 你的控制台点不动了")
para("看控制台右上角的提示：")
bullet("说明是你这台设备的网络问题。换一台备用设备打开同一个控制台链接，直接接管，分数全在。",
       "显示「本机掉线 · 服务器正常」——")
bullet("跳到第四级。", "显示「服务器不可达」——")
say("稍等一下，我换个设备。")

h2("三级 · 全场同时掉线")
para("所有人（包括大屏）一起断了。系统会自己重启并恢复所有人的分数。")
warn("等 10 秒。不要做任何操作。")
say("系统打个盹，马上回来，大家的分数一分都不会少。"
    "正好趁这个空当，让我们再次把掌声送给今天的新郎新娘 ——")
para("恢复之后会停在上一题的结果页。由你决定：按「换一道备用题」把刚才那题重来，"
     "或者直接按「下一题」，把那题作废继续往下。")

h2("四级 · 彻底连不上")
para("判断标准很简单：", )
warn("等满 30 秒还是打不开，就放弃系统，切换到纸质玩法。")
say("看来今天的网络不太配合。没关系，我们换个更热闹的玩法 —— "
    "接下来这几题我念出来，谁先举手谁回答，答对的有奖！")
para("你手上会有一份纸质题卡（含答案），并且会安排一位记分员带着纸笔。"
     "这一级不需要任何技术支持，你一个人就能把场子撑住。")

# ──────────────────────────────────────────────
h1("八、绝对不要做的四件事")

bullet("正确答案就在上面。", "不要把控制台屏幕转向观众 —— ")
bullet("倒计时会立刻开始走，宾客要一边听你念一边看时间掉。", "不要先按「下一题」再念题 —— ")
bullet("这两个按钮按下去撤不回来，会影响全场的分数。", "不要随手点「跳过本题」和「换一道备用题」—— ")
bullet("那会结束整场游戏，之后不能再让人入场，也不能再答题。系统会问你一次，看清楚再确认。",
       "不要在还没答完 13 题时点「公布最终排名」—— ")

# ──────────────────────────────────────────────
doc.add_page_break()
title("速查卡", "建议单独打印一张，放在手边")

h1("每题四步")
para("① 念完题和四个选项　→　② 按「下一题」　→　③ 等 20 秒，唱分　→　④ 按「下一题」")

h1("唱分模板")
say("这一题，＿＿位来宾作答，＿＿位答对，正确答案是＿！目前第一名是「＿＿＿＿」，＿＿＿＿分。")

h1("三句救命的话")
table(
    ["情况", "直接念"],
    [
        ["大屏黑了", "「大屏这边稍微调整一下，大家看手机就可以，我们继续。」"],
        ["全场掉线", "「系统打个盹，马上回来，大家的分数一分都不会少。正好趁这个空当，"
                    "让我们再次把掌声送给今天的新郎新娘 ——」"],
        ["彻底连不上", "「看来今天的网络不太配合。我们换个更热闹的玩法 —— "
                     "接下来这几题我念出来，谁先举手谁回答，答对的有奖！」"],
    ],
    [2.6, 11.6],
)

h1("四个数字")
para("13 道题　·　每题 20 秒　·　延长一次 10 秒（该时段内作答 0 分）　·　全场约 15 分钟")

h1("记住这一条")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = _cn(p.add_run("题目在每个人手机上都有。"), "楷体")
r.font.size = Pt(15)
r.bold = True
r.font.color.rgb = SEAL
p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = _cn(p2.add_run("大屏只是放大给大家看 —— 它出问题，游戏不会停。"), "楷体")
r2.font.size = Pt(13)
r2.font.color.rgb = INK2

import os
os.chdir(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
out = "docs/主持人培训手册.docx"
doc.save(out)
print("已生成", out)
